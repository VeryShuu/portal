"""Integration: экспорт прогресса в xlsx (этап 2) — раздача методисту,
403 для не-методиста, содержимое соответствует прогрессу."""

from __future__ import annotations

import contextlib
import io
import uuid
from types import SimpleNamespace

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.database import AsyncSessionLocal
from app.models.learning import LearningCourse

pytestmark = pytest.mark.asyncio


@pytest_asyncio.fixture(autouse=True)
async def _cleanup_learning_rows(redis_client):
    yield

    async with AsyncSessionLocal() as db:
        await db.execute(text("DELETE FROM email_outbox WHERE kind = 'learning'"))
        await db.execute(text("DELETE FROM learning_item_progress"))
        await db.execute(text("DELETE FROM learning_test_attempts"))
        await db.execute(text("DELETE FROM learning_question_options"))
        await db.execute(text("DELETE FROM learning_tests"))
        await db.execute(text("DELETE FROM learning_course_items"))
        await db.execute(text("DELETE FROM learning_courses"))
        await db.execute(text("DELETE FROM learning_course_participants"))
        await db.execute(text("DELETE FROM learning_admins"))
        await db.execute(text("DELETE FROM learning_accounts WHERE email LIKE '%@example.com'"))
        await db.commit()
    for prefix in ("learning_session:", "learning_sessions:"):
        for key in await redis_client.keys(f"{prefix}*"):
            await redis_client.delete(key)
    for key in await redis_client.keys("session:sid-pexp-*"):
        await redis_client.delete(key)


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("ADMIN_EMAIL", "")
    monkeypatch.setenv("ADMIN_PASSWORD", "")
    import importlib

    import app.main as main_mod

    importlib.reload(main_mod)
    return main_mod.app


@pytest_asyncio.fixture
async def live_limiter(redis_client):
    from fastapi_limiter import FastAPILimiter
    from fastapi_limiter.depends import RateLimiter

    import tests.conftest as _root_conftest

    saved_call = RateLimiter.__call__
    if getattr(_root_conftest, "_real_rate_limiter_call", None) is not None:
        RateLimiter.__call__ = _root_conftest._real_rate_limiter_call  # type: ignore[method-assign]

    await FastAPILimiter.init(redis_client)
    try:
        yield redis_client
    finally:
        RateLimiter.__call__ = saved_call  # type: ignore[method-assign]
        with contextlib.suppress(Exception):
            await FastAPILimiter.close()


def _enable_module(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.core import modules_config

    async def fake_load(redis):
        return SimpleNamespace(learning=SimpleNamespace(enabled=True))

    monkeypatch.setattr(modules_config, "load_modules_shared", fake_load)


async def _seed_progress_course() -> tuple[uuid.UUID, uuid.UUID, uuid.UUID]:
    """Курс с материалом, прошедший сотрудник и методист."""
    from datetime import UTC, datetime

    from app.models.learning import (
        LearningAdmin,
        LearningCourseItem,
        LearningCourseParticipant,
        LearningItemProgress,
    )
    from app.models.user import User

    async with AsyncSessionLocal() as db:
        course = LearningCourse(
            slug=f"pexp-{uuid.uuid4().hex[:8]}",
            title="Курс экспорта",
            status="published",
            published_at=datetime.now(UTC),
        )
        db.add(course)
        await db.flush()
        material = LearningCourseItem(
            course_id=course.id, type="material", title="Материал", sort_order=0
        )
        db.add(material)
        admin = User(
            email=f"pexp-admin-{uuid.uuid4().hex[:8]}@example.com",
            full_name="Методист Экспорт",
            role="reader",
            auth_source="keycloak",
            current_status="working",
        )
        staff = User(
            email=f"pexp-staff-{uuid.uuid4().hex[:8]}@example.com",
            full_name="Сотрудник Экспорт",
            role="reader",
            auth_source="keycloak",
            current_status="working",
        )
        db.add_all([admin, staff])
        await db.flush()
        db.add(LearningAdmin(user_id=admin.id))
        db.add(
            LearningCourseParticipant(
                course_id=course.id,
                user_id=staff.id,
                display_name=staff.full_name,
                email=staff.email,
            )
        )
        db.add(LearningItemProgress(item_id=material.id, user_id=staff.id))
        await db.commit()
        return course.id, admin.id, staff.id


def _client(app, sid: str) -> AsyncClient:
    return AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
        headers={"Origin": "http://test", "Cookie": f"portal_session={sid}"},
        follow_redirects=False,
    )


class TestProgressExport:
    async def test_export_contains_progress(self, app, redis_client, live_limiter, monkeypatch):
        _enable_module(monkeypatch)
        from app.services.session import save_session

        course_id, admin_uid, _staff_uid = await _seed_progress_course()
        await save_session(
            redis_client, "sid-pexp-admin", {"user_id": str(admin_uid), "auth_source": "local"}
        )

        app.state.redis = redis_client
        async with _client(app, "sid-pexp-admin") as ac:
            r = await ac.get(f"/api/v1/learning/admin/courses/{course_id}/progress/export")
            assert r.status_code == 200, r.text
            assert r.content[:2] == b"PK"
            assert "attachment" in r.headers.get("content-disposition", "")

            wb = load_workbook(io.BytesIO(r.content))
            ws = wb.active
            table = list(ws.iter_rows(values_only=True))
            assert table[0][0] == "Участник"
            assert len(table) == 2  # заголовок + одна строка
            assert table[1][0] == "Сотрудник Экспорт"
            assert table[1][1] == "Сотрудник"
            assert table[1][4] == 1 and table[1][5] == 1

        await redis_client.delete("session:sid-pexp-admin")

    async def test_requires_learning_admin(self, app, redis_client, live_limiter, monkeypatch):
        _enable_module(monkeypatch)
        from app.services.session import save_session

        course_id, _admin_uid, staff_uid = await _seed_progress_course()
        await save_session(
            redis_client, "sid-pexp-admin", {"user_id": str(staff_uid), "auth_source": "local"}
        )

        app.state.redis = redis_client
        async with _client(app, "sid-pexp-admin") as ac:
            r = await ac.get(f"/api/v1/learning/admin/courses/{course_id}/progress/export")
            assert r.status_code == 403
