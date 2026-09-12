"""Unit-тесты экспорта прогресса (этап 2): колонки = таблице прогресса,
строки сортированы по зачислению, сертификат — «да»/пусто."""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services.learning import courses_service as cs
from app.services.learning.progress_export import build_progress_xlsx


def _row(
    name: str,
    kind: str,
    email: str,
    enrolled_at: datetime,
    completed: int,
    *,
    certificate: bool = False,
) -> cs.ProgressRow:
    item_a, item_b = uuid.uuid4(), uuid.uuid4()
    return cs.ProgressRow(
        participant_id=uuid.uuid4(),
        kind=kind,
        display_name=name,
        email=email,
        enrolled_at=enrolled_at,
        completed_items=frozenset({item_a} if completed else set()),
        passed_tests=frozenset({item_b} if completed > 1 else set()),
        has_certificate=certificate,
    )


class TestBuildProgressXlsx:
    def test_headers_and_rows(self):
        rows = [
            _row(
                "Иванов И.",
                "staff",
                "ivanov@x",
                datetime(2026, 8, 1, tzinfo=UTC),
                completed=2,
                certificate=True,
            ),
            _row(
                "Петров П.",
                "external",
                "petrov@x",
                datetime(2026, 8, 2, tzinfo=UTC),
                completed=0,
            ),
        ]
        data = build_progress_xlsx(course_title="Курс", total_items=3, rows=rows)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        table = list(ws.iter_rows(values_only=True))
        assert table[0] == (
            "Участник",
            "Тип",
            "Email",
            "Зачислен",
            "Пройдено элементов",
            "Всего элементов",
            "Сертификат",
        )
        assert [r[0] for r in table[1:]] == ["Иванов И.", "Петров П."]
        assert table[1][1] == "Сотрудник" and table[2][1] == "Внешняя учётка"
        assert table[1][4] == 2 and table[1][5] == 3
        assert table[1][6] == "да" and table[2][6] in ("", None)
