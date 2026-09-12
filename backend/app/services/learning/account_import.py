"""XLSX-импорт внешних учёток обучения (§5.1 ТЗ). Паролей нет (passwordless,
миграция 113): учётка = email + ФИО; вход — кодом из письма при первом заходе."""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.learning import LearningAccount
from app.schemas.learning import LearningAccountCreate

MAX_IMPORT_ROWS = 1000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 500

_TEMPLATE_HEADERS = ("Фамилия Имя", "email", "отдел", "должность")
_HEADER_ALIASES: dict[str, set[str]] = {
    "full_name": {"фамилия имя", "имя фамилия", "фио", "full_name", "full name"},
    "email": {"email", "e-mail", "электронная почта"},
    "department": {"отдел", "подразделение", "department"},
    "position": {"должность", "position", "job title"},
}


class ImportFormatError(ValueError):
    """Файл целиком не может быть разобран как допустимый шаблон."""


@dataclass(frozen=True, slots=True)
class ImportRow:
    row_number: int
    account: LearningAccountCreate


@dataclass(frozen=True, slots=True)
class ImportRowError:
    row_number: int
    message: str


@dataclass(frozen=True, slots=True)
class ParsedImport:
    rows: list[ImportRow]
    errors: list[ImportRowError]


@dataclass(frozen=True, slots=True)
class ImportResult:
    created: int
    skipped_duplicates: int
    errors: list[ImportRowError]


def _cell_text(value: object) -> str:
    """Excel-типы намеренно приводятся к строке до бизнес-валидации."""
    return "" if value is None else str(value).strip()


def _normalise_header(value: object) -> str:
    return " ".join(_cell_text(value).lower().replace("ё", "е").split())


def _header_map(values: tuple[object, ...]) -> dict[str, int]:
    result: dict[str, int] = {}
    for column, raw in enumerate(values):
        header = _normalise_header(raw)
        for field, aliases in _HEADER_ALIASES.items():
            if header not in aliases:
                continue
            if field in result:
                raise ImportFormatError(f"Колонка «{_cell_text(raw)}» указана повторно")
            result[field] = column
            break
    missing = [name for name in ("full_name", "email") if name not in result]
    if missing:
        raise ImportFormatError("Нужны обязательные колонки «Фамилия Имя» и «email»")
    return result


def _validation_message(exc: ValidationError) -> str:
    error = exc.errors(include_url=False)[0]
    field = str(error.get("loc", ("строка",))[0])
    labels = {
        "email": "email",
        "full_name": "Фамилия Имя",
        "department": "отдел",
        "position": "должность",
    }
    error_type = str(error.get("type", ""))
    if error_type in {"missing", "string_type", "string_too_short"}:
        message = "обязательное поле"
    elif error_type == "string_too_long":
        message = "значение превышает допустимую длину"
    elif error_type == "value_error":
        context_error = error.get("ctx", {}).get("error")
        message = str(context_error) if context_error else "некорректное значение"
    else:
        message = "некорректное значение"
    return f"{labels.get(field, field)}: {message}"


def _row_value(values: tuple[object, ...], columns: dict[str, int], field: str) -> str | None:
    column = columns.get(field)
    if column is None or column >= len(values):
        return None
    text = _cell_text(values[column])
    return text or None


def parse_xlsx(data: bytes) -> ParsedImport:
    """Разобрать первый лист xlsx; ошибочные строки вернуть в отчёте."""
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise ImportFormatError("В xlsx слишком много внутренних файлов")
            if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
                raise ImportFormatError("Распакованный xlsx превышает лимит 50 МБ")
    except BadZipFile as exc:
        raise ImportFormatError("Файл не является корректным xlsx") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ImportFormatError("Файл не является корректным xlsx") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFormatError("В файле нет листов")
        iterator = sheet.iter_rows(values_only=True)
        header_row_number = 0
        header_values: tuple[object, ...] | None = None
        for row_number, values in enumerate(iterator, start=1):
            if any(_cell_text(value) for value in values):
                header_row_number = row_number
                header_values = values
                break
        if header_values is None:
            raise ImportFormatError("Файл пуст")

        columns = _header_map(header_values)
        rows: list[ImportRow] = []
        errors: list[ImportRowError] = []
        data_rows = 0
        for row_number, values in enumerate(iterator, start=header_row_number + 1):
            if not any(_cell_text(value) for value in values):
                continue
            data_rows += 1
            if data_rows > MAX_IMPORT_ROWS:
                raise ImportFormatError(
                    f"В одном файле допускается не более {MAX_IMPORT_ROWS} строк"
                )

            try:
                account = LearningAccountCreate.model_validate(
                    {
                        "full_name": _row_value(values, columns, "full_name"),
                        "email": _row_value(values, columns, "email"),
                        "department": _row_value(values, columns, "department"),
                        "position": _row_value(values, columns, "position"),
                    }
                )
            except ValidationError as exc:
                errors.append(ImportRowError(row_number, _validation_message(exc)))
                continue
            rows.append(ImportRow(row_number, account))
        return ParsedImport(rows=rows, errors=errors)
    finally:
        workbook.close()


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Учётки"
    sheet.append(_TEMPLATE_HEADERS)
    sheet.append(("Иванов Иван", "ivanov@example.ru", "Продажи", "Менеджер"))
    fill = PatternFill("solid", fgColor="305496")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip("ABCD", (34, 34, 30, 30), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


async def import_accounts(db: AsyncSession, parsed: ParsedImport) -> ImportResult:
    """Создать валидные уникальные строки одной транзакцией. Письма при импорте
    не отправляются: паролей больше нет, а сам факт создания учётки действия
    не требует — человек зайдёт кодом, когда до него дойдёт задача."""
    unique: list[ImportRow] = []
    seen: set[str] = set()
    skipped = 0
    for row in parsed.rows:
        email = row.account.email.lower()
        if email in seen:
            skipped += 1
            continue
        seen.add(email)
        unique.append(row)

    occupied: set[str] = set()
    if seen:
        occupied = {
            str(email)
            for email in (
                await db.execute(
                    select(func.lower(LearningAccount.email)).where(
                        func.lower(LearningAccount.email).in_(seen)
                    )
                )
            ).scalars()
        }
    pending = [row for row in unique if row.account.email.lower() not in occupied]
    skipped += len(unique) - len(pending)

    values = [
        {
            "id": uuid.uuid4(),
            "email": row.account.email,
            "full_name": row.account.full_name,
            "department": row.account.department,
            "position": row.account.position,
        }
        for row in pending
    ]
    created = 0
    if values:
        statement = (
            pg_insert(LearningAccount)
            .values(values)
            .on_conflict_do_nothing(
                index_elements=[func.lower(LearningAccount.email)],
                index_where=LearningAccount.deleted_at.is_(None),
            )
            .returning(LearningAccount.id)
        )
        created = len((await db.execute(statement)).all())
        skipped += len(values) - created
    return ImportResult(
        created=created,
        skipped_duplicates=skipped,
        errors=parsed.errors,
    )
