"""Экспорт прогресса курса в xlsx (этап 2; ТЗ §13, §15).

Колонки выгрузки = колонкам таблицы прогресса (§13: файл читается без
объяснений); заголовки те же термины, что в админ-UI. Экспорт — обычная
генерация openpyxl (уже в зависимостях после импорта учёток/вопросов).
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.learning import courses_service as cs

EXPORT_HEADERS = (
    "Участник",
    "Тип",
    "Email",
    "Зачислен",
    "Пройдено элементов",
    "Всего элементов",
    "Сертификат",
)


def _kind_label(kind: str) -> str:
    return "Сотрудник" if kind == "staff" else "Внешняя учётка"


def _cert_label(has_certificate: bool) -> str:
    return "да" if has_certificate else ""


def build_progress_xlsx(
    *,
    course_title: str,
    total_items: int,
    rows: list[cs.ProgressRow],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Прогресс"
    sheet.append(EXPORT_HEADERS)

    fill = PatternFill("solid", fgColor="305496")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for r in sorted(rows, key=lambda x: x.enrolled_at):
        sheet.append(
            (
                r.display_name,
                _kind_label(r.kind),
                r.email,
                r.enrolled_at.strftime("%d.%m.%Y %H:%M"),
                len(r.completed_items | r.passed_tests),
                total_items,
                _cert_label(r.has_certificate),
            )
        )

    for column, width in zip("ABCDEFG", (36, 18, 34, 20, 18, 16, 14), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
